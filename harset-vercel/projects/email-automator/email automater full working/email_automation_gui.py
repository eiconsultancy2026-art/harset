import os,re,csv,ssl,smtplib,tkinter as tk
from tkinter import ttk,filedialog,messagebox
from datetime import datetime
from email.message import EmailMessage
from openpyxl import load_workbook

MASTER={'s.no','name of trust','email'}
TEMPLATE={'sl. no.','pre acknowledgement number','id code','unique registration number (urn)','name of donor','address of donor','donation type','mode of receipt','amount of donation (indian rupees)'}
DEFAULT_SUBJECT='Request for Donation Details'
DEFAULT_BODY="""Dear {trust_name},

Greetings from [Company/Firm Name].

Please find attached the Excel format for providing the required donation details.

We request you to kindly fill in the required information in the attached Excel sheet and send the completed sheet back to us at your earliest convenience.

Please ensure that all the applicable details are entered accurately and completely.

If you have any questions or require any clarification regarding the information to be provided, please feel free to contact us.

Thank you for your cooperation.

Regards,
[Your Name]
[Company/Firm Name]
[Contact Number]"""
EMAIL_RE=re.compile(r'^[^\s@]+@[^\s@]+\.[^\s@]+$')
def norm(x): return re.sub(r'\s+',' ',str(x or '').strip().lower())
def rows(path):
    wb=load_workbook(path,data_only=True,read_only=True); r=list(wb.active.iter_rows(values_only=True)); wb.close(); return r
def header(rs,wanted):
    wanted=set(wanted)
    for i,r in enumerate(rs[:20]):
        vals={norm(x) for x in r if x is not None}
        if wanted.issubset(vals): return i,r
    return None,None
def load_master(path):
    rs=rows(path); i,h=header(rs,MASTER)
    if i is None: raise ValueError('Trust master must contain S.No, Name of Trust and email.')
    hs=[str(x).strip() if x is not None else '' for x in h]; ix={norm(x):j for j,x in enumerate(hs)}; out=[]; bad=[]
    for rn,r in enumerate(rs[i+1:],i+2):
        v=[str(x).strip() if x is not None else '' for x in r]
        if not any(v): continue
        get=lambda k: v[ix[k]] if ix[k]<len(v) else ''
        sno,trust,email=get('s.no'),get('name of trust'),get('email')
        if trust and EMAIL_RE.match(email): out.append((sno,trust,email))
        else: bad.append(rn)
    if not out: raise ValueError('No valid clients found in the trust master.')
    return out,bad
def validate_template(path):
    if header(rows(path),TEMPLATE)[0] is None: raise ValueError('Client template does not contain the required donation-detail columns.')
def send_all(sender,pw,subject,body,clients,attach,log,progress):
    os.makedirs(os.path.dirname(log),exist_ok=True); new=not os.path.exists(log); sent=failed=0
    with open(log,'a',newline='',encoding='utf-8') as f:
        w=csv.writer(f)
        if new:w.writerow(['timestamp','sender_email','sno','trust_name','recipient_email','subject','attachment','status','detail'])
        with smtplib.SMTP_SSL('smtp.gmail.com',465,context=ssl.create_default_context(),timeout=30) as s:
            s.login(sender,pw)
            with open(attach,'rb') as a:data=a.read()
            for n,(sno,trust,email) in enumerate(clients,1):
                try:
                    m=EmailMessage();m['From']=sender;m['To']=email;m['Subject']=subject;m.set_content(body.replace('{trust_name}',trust));m.add_attachment(data,maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=os.path.basename(attach));s.send_message(m);status='SENT';detail='';sent+=1
                except Exception as e:status='FAILED';detail=str(e);failed+=1
                w.writerow([datetime.now().strftime('%Y-%m-%d %H:%M:%S'),sender,sno,trust,email,subject,os.path.basename(attach),status,detail]);f.flush();progress(n,len(clients),status,trust)
    return sent,failed
class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title('Client Excel Request Email Automater');self.geometry('1000x760');self.minsize(850,650);self.clients=[];self.master_path='';self.template='';self.build()
    def build(self):
        p={'padx':18,'pady':12};self.configure(**p);ttk.Label(self,text='Client Excel Request Email Automater',font=('Segoe UI',18,'bold')).pack(anchor='w');ttk.Label(self,text='Upload Excel files, edit the default email, preview recipients, then send.').pack(anchor='w',pady=(0,12))
        f=ttk.LabelFrame(self,text='1. Excel Files',padding=10);f.pack(fill='x');self.mv=tk.StringVar();self.tv=tk.StringVar();self.file(f,'Trust Master (S.No | Name of Trust | email)',self.mv,self.pickm);self.file(f,'Client Excel Template (donation details)',self.tv,self.pickt)
        c=ttk.LabelFrame(self,text='2. Email Settings — editable',padding=10);c.pack(fill='x',pady=10);c.columnconfigure(1,weight=1);c.columnconfigure(3,weight=1)
        ttk.Label(c,text='Sender Gmail').grid(row=0,column=0,sticky='w');self.sender=tk.StringVar();ttk.Entry(c,textvariable=self.sender).grid(row=0,column=1,sticky='ew',padx=8)
        ttk.Label(c,text='Gmail App Password').grid(row=0,column=2,sticky='w');self.pw=tk.StringVar();ttk.Entry(c,textvariable=self.pw,show='*').grid(row=0,column=3,sticky='ew',padx=8)
        ttk.Label(c,text='Subject').grid(row=1,column=0,sticky='w',pady=6);self.sub=tk.StringVar(value=DEFAULT_SUBJECT);ttk.Entry(c,textvariable=self.sub).grid(row=1,column=1,columnspan=3,sticky='ew',padx=8)
        ttk.Label(c,text='Email Body').grid(row=2,column=0,sticky='nw');self.body=tk.Text(c,height=12,wrap='word');self.body.grid(row=2,column=1,columnspan=3,sticky='ew',padx=8,pady=6);self.body.insert('1.0',DEFAULT_BODY);ttk.Label(c,text='Use {trust_name} wherever you want the trust name inserted automatically.').grid(row=3,column=1,columnspan=3,sticky='w')
        a=ttk.Frame(self);a.pack(fill='x',pady=6);self.loadb=ttk.Button(a,text='Load & Preview Recipients',command=self.load);self.loadb.pack(side='left');self.sendb=ttk.Button(a,text='SEND ALL',command=self.send,state='disabled');self.sendb.pack(side='right');self.status=tk.StringVar(value='Select both Excel files.');ttk.Label(self,textvariable=self.status).pack(anchor='w');self.prog=ttk.Progressbar(self);self.prog.pack(fill='x',pady=5)
        box=ttk.LabelFrame(self,text='3. Recipient Preview',padding=6);box.pack(fill='both',expand=True);self.tree=ttk.Treeview(box,columns=('sno','trust','email'),show='headings');
        for col,title,w in [('sno','S.No',70),('trust','Name of Trust',320),('email','Email',420)]:self.tree.heading(col,text=title);self.tree.column(col,width=w)
        self.tree.pack(side='left',fill='both',expand=True);sb=ttk.Scrollbar(box,command=self.tree.yview);sb.pack(side='right',fill='y');self.tree.configure(yscrollcommand=sb.set)
    def file(self,p,label,var,cmd):
        r=ttk.Frame(p);r.pack(fill='x',pady=4);ttk.Label(r,text=label,width=40).pack(side='left');ttk.Entry(r,textvariable=var).pack(side='left',fill='x',expand=True,padx=8);ttk.Button(r,text='Browse',command=cmd).pack(side='right')
    def pickm(self):
        p=filedialog.askopenfilename(filetypes=[('Excel files','*.xlsx')]);
        if p:self.master_path=p;self.mv.set(p)
    def pickt(self):
        p=filedialog.askopenfilename(filetypes=[('Excel files','*.xlsx')]);
        if p:self.template=p;self.tv.set(p)
    def load(self):
        try:
            if not self.master_path or not self.template:raise ValueError('Select both Excel files.')
            validate_template(self.template);self.clients,bad=load_master(self.master_path)
            self.tree.delete(*self.tree.get_children())
            for x in self.clients:self.tree.insert('', 'end',values=x)
            self.sendb.config(state='normal');self.status.set(f'{len(self.clients)} recipients loaded. {len(bad)} invalid rows skipped.')
        except Exception as e:messagebox.showerror('Validation Error',str(e))
    def send(self):
        sender=self.sender.get().strip();pw=self.pw.get();sub=self.sub.get().strip() or DEFAULT_SUBJECT;body=self.body.get('1.0','end').strip()
        if not EMAIL_RE.match(sender):return messagebox.showerror('Error','Enter a valid sender Gmail address.')
        if not pw:return messagebox.showerror('Error','Enter your Gmail App Password.')
        if not body:return messagebox.showerror('Error','Email body cannot be empty.')
        if not messagebox.askyesno('Confirm Send',f'Send {len(self.clients)} emails?\n\nSubject: {sub}\nAttachment: {os.path.basename(self.template)}'):return
        self.sendb.config(state='disabled');self.loadb.config(state='disabled')
        try:
            log=os.path.join(os.path.dirname(os.path.abspath(__file__)),'output','delivery_log.csv');sent,failed=send_all(sender,pw,sub,body,self.clients,self.template,log,self.progress)
            messagebox.showinfo('Complete',f'Sent: {sent}\nFailed: {failed}\n\nLog: output\\delivery_log.csv');self.status.set(f'Completed: {sent} sent, {failed} failed.')
        except smtplib.SMTPAuthenticationError:messagebox.showerror('Gmail Login Failed','Gmail rejected the login. Use a 16-character Gmail App Password, not your normal password.')
        except Exception as e:messagebox.showerror('Send Error',str(e))
        finally:self.sendb.config(state='normal');self.loadb.config(state='normal');self.pw.set('')
    def progress(self,n,total,status,trust):self.prog['maximum']=total;self.prog['value']=n;self.status.set(f'{n}/{total}: {status} - {trust}');self.update_idletasks()
if __name__=='__main__':App().mainloop()
