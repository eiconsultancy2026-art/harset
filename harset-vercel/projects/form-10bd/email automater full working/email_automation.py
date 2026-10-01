#!/usr/bin/env python3
"""
CLIENT EXCEL REQUEST EMAIL AUTOMATER
====================================

Reads a Trust/Client master Excel file containing:
    S.No | Name of Trust | email

and sends the SAME client-input Excel template as an attachment to every
valid client. The email greeting is personalized with the Trust name.

No Word/DOCX generation is used in this version.

Security:
- Gmail App Password is requested interactively and never saved.
- App Password is never written to logs or files.
"""

import csv
import glob
import os
import re
import smtplib
import ssl
import sys
import time
from datetime import datetime, timedelta
from email.message import EmailMessage
from getpass import getpass

try:
    from openpyxl import load_workbook
except ImportError:
    print("Missing dependency 'openpyxl'. Run: pip install openpyxl")
    sys.exit(1)

INPUT_DIR = "input"
OUTPUT_DIR = "output"
LOG_FILE = os.path.join(OUTPUT_DIR, "delivery_log.csv")

GMAIL_SMTP_HOST = "smtp.gmail.com"
GMAIL_SMTP_PORT = 465

REQUIRED_MASTER_HEADERS = ["S.No", "Name of Trust", "email"]
REQUIRED_TEMPLATE_HEADERS = [
    "Sl. No.",
    "Pre Acknowledgement Number",
    "ID Code",
    "Unique Registration Number (URN)",
    "Name of donor",
    "Address of donor",
    "Donation Type",
    "Mode of receipt",
    "Amount of donation (Indian rupees)",
]

DEFAULT_SUBJECT = "Request for Donation Details"
DEFAULT_BODY = """Dear {trust_name},

Please find attached the Excel format for providing the required donation details.

Kindly fill in the required details in the attached Excel sheet and send the completed sheet back to us at your earliest convenience.

Please ensure that all applicable fields are filled in accurately.

Regards,
[Your Name]
[Company/Firm Name]
"""

EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


def section(title):
    print("\n" + "=" * 68)
    print(title)
    print("=" * 68 + "\n")


def normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def verify_file_accessible(path, label):
    path = os.path.abspath(path)
    try:
        with open(path, "rb") as f:
            f.read(8)
    except FileNotFoundError:
        print(f"ERROR: {label} was not found:\n{path}")
        sys.exit(1)
    except PermissionError:
        print(f"ERROR: Could not open {label}. It may be locked by Excel:\n{path}")
        print("Close the file in Excel and run the program again.")
        sys.exit(1)


def read_sheet_rows(path):
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        return rows
    except PermissionError:
        print(f"ERROR: Could not open Excel file because it may be locked:\n{path}")
        print("Close the file in Excel and try again.")
        sys.exit(1)
    except Exception as exc:
        print(f"ERROR: Could not read Excel file:\n{path}")
        print(f"Details: {exc}")
        sys.exit(1)


def find_header_row(rows, required_headers):
    wanted = {normalize_header(h) for h in required_headers}
    for row_index, row in enumerate(rows[:15]):
        actual = {normalize_header(cell) for cell in row if cell is not None}
        if wanted.issubset(actual):
            return row_index
    return None


def headers_from_row(row):
    return [str(cell).strip() if cell is not None else "" for cell in row]


def classify_excel_files():
    """Find the two Excel files by their column structure, not filename."""
    if not os.path.isdir(INPUT_DIR):
        print(f"ERROR: Input folder '{INPUT_DIR}' does not exist.")
        print("Create it and place the two Excel files inside it.")
        sys.exit(1)

    files = [
        f for f in glob.glob(os.path.join(INPUT_DIR, "*.xlsx"))
        if not os.path.basename(f).startswith("~$")
    ]

    if len(files) < 2:
        print("ERROR: Two .xlsx files are required in input/:")
        print("  1. Trust master: S.No, Name of Trust, email")
        print("  2. Client template: the donation-detail columns")
        sys.exit(1)

    master = None
    template = None

    for path in files:
        verify_file_accessible(path, "Excel file")
        rows = read_sheet_rows(path)
        if not rows:
            continue

        if find_header_row(rows, REQUIRED_MASTER_HEADERS) is not None:
            master = path
        if find_header_row(rows, REQUIRED_TEMPLATE_HEADERS) is not None:
            template = path

    if master is None:
        print("ERROR: Could not identify the Trust master Excel.")
        print("It must contain: S.No, Name of Trust, email")
        sys.exit(1)

    if template is None:
        print("ERROR: Could not identify the client-input Excel template.")
        print("It must contain these columns:")
        for h in REQUIRED_TEMPLATE_HEADERS:
            print(f"  - {h}")
        sys.exit(1)

    if os.path.abspath(master) == os.path.abspath(template):
        print("ERROR: The same Excel file cannot be both master and template.")
        sys.exit(1)

    print("Trust master detected:")
    print(f"  {os.path.basename(master)}")
    print("Client-input template detected:")
    print(f"  {os.path.basename(template)}")
    return master, template


def load_recipients(master_path):
    rows = read_sheet_rows(master_path)
    header_row_idx = find_header_row(rows, REQUIRED_MASTER_HEADERS)
    if header_row_idx is None:
        print("ERROR: Invalid trust master Excel.")
        sys.exit(1)

    headers = headers_from_row(rows[header_row_idx])
    normalized_to_real = {normalize_header(h): h for h in headers if h}

    sno_h = normalized_to_real[normalize_header("S.No")]
    trust_h = normalized_to_real[normalize_header("Name of Trust")]
    email_h = normalized_to_real[normalize_header("email")]

    recipients = []
    invalid_rows = []

    for excel_row, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        record = {}
        for idx, header in enumerate(headers):
            if header:
                value = row[idx] if idx < len(row) else None
                record[header] = "" if value is None else str(value).strip()

        sno = record.get(sno_h, "")
        trust = record.get(trust_h, "").strip()
        email = record.get(email_h, "").strip()

        if not trust or not email or not EMAIL_RE.match(email):
            invalid_rows.append((excel_row, sno, trust, email))
            continue

        recipients.append({"sno": sno, "trust_name": trust, "email": email})

    if not recipients:
        print("ERROR: No valid clients were found in the trust master Excel.")
        sys.exit(1)

    print(f"Valid clients found: {len(recipients)}")
    if invalid_rows:
        print(f"Rows skipped because Trust name/email was invalid: {len(invalid_rows)}")
        for row in invalid_rows[:10]:
            print(f"  Excel row {row[0]}: S.No={row[1]!r}, Trust={row[2]!r}, Email={row[3]!r}")
        if len(invalid_rows) > 10:
            print(f"  ...and {len(invalid_rows) - 10} more")

    return recipients


def validate_template(template_path):
    rows = read_sheet_rows(template_path)
    header_row_idx = find_header_row(rows, REQUIRED_TEMPLATE_HEADERS)
    if header_row_idx is None:
        print("ERROR: Client-input template does not match the required format.")
        sys.exit(1)
    print(f"Template validation: OK (header row {header_row_idx + 1})")


def prompt_gmail_credentials():
    section("GMAIL SENDER CONFIGURATION")
    sender_email = input("Enter sender Gmail address: ").strip()
    sender_password = getpass("Enter Gmail App Password: ")
    return sender_email, sender_password


def test_gmail_login(sender_email, sender_password):
    context = ssl.create_default_context()
    try:
        with smtplib.SMTP_SSL(GMAIL_SMTP_HOST, GMAIL_SMTP_PORT, context=context, timeout=20) as server:
            server.login(sender_email, sender_password)
        return True, ""
    except smtplib.SMTPAuthenticationError as exc:
        return False, f"Gmail authentication failed: {exc}"
    except (OSError, smtplib.SMTPException) as exc:
        return False, f"Could not connect to Gmail: {exc}"


def get_authenticated_gmail_credentials():
    while True:
        sender_email, sender_password = prompt_gmail_credentials()
        print("\nTesting Gmail connection...")
        ok, detail = test_gmail_login(sender_email, sender_password)
        if ok:
            print("Gmail authentication successful.")
            return sender_email, sender_password

        print("\n" + detail)
        print("Check that:")
        print("  1. The Gmail address is correct.")
        print("  2. 2-Step Verification is enabled.")
        print("  3. You are using a Gmail App Password, not your normal password.")
        choice = input("Type R to retry or X to exit: ").strip().lower()
        if choice != "r":
            sys.exit(1)


def prompt_email_configuration():
    section("EMAIL CONFIGURATION")

    subject = input(f"Email subject [{DEFAULT_SUBJECT}]: ").strip() or DEFAULT_SUBJECT

    print("Enter the email body.")
    print("Use {trust_name} where you want the Trust name to appear.")
    print("Press Enter twice on an empty line when finished.")
    print("\nExample:\n" + DEFAULT_BODY)
    print("---")

    lines = []
    while True:
        line = input()
        if line == "" and lines and lines[-1] == "":
            break
        lines.append(line)

    body = "\n".join(lines).strip() if lines else DEFAULT_BODY.strip()
    return subject, body


def delivery_confirmation(sender_email, recipients, template_path, subject):
    section("DELIVERY CONFIRMATION")
    print(f"Sender:     {sender_email}")
    print(f"Recipients: {len(recipients)}")
    print(f"Attachment: {os.path.basename(template_path)}")
    print(f"Subject:    {subject}")
    print("\nFirst recipients:")
    for r in recipients[:10]:
        print(f"  {r['sno']} | {r['trust_name']} | {r['email']}")
    if len(recipients) > 10:
        print(f"  ...and {len(recipients) - 10} more")

    choice = input("\nType SEND to send all emails, or X to cancel: ").strip()
    return choice == "SEND"


def build_email(sender_email, recipient, subject, body_template, attachment_path):
    msg = EmailMessage()
    msg["From"] = sender_email
    msg["To"] = recipient["email"]
    msg["Subject"] = subject

    body = body_template.replace("{trust_name}", recipient["trust_name"])
    msg.set_content(body)

    with open(attachment_path, "rb") as f:
        data = f.read()

    msg.add_attachment(
        data,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=os.path.basename(attachment_path),
    )
    return msg


def init_log():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    is_new = not os.path.exists(LOG_FILE)
    f = open(LOG_FILE, "a", newline="", encoding="utf-8")
    writer = csv.writer(f)
    if is_new:
        writer.writerow([
            "timestamp",
            "sender_email",
            "sno",
            "trust_name",
            "recipient_email",
            "subject",
            "attachment",
            "status",
            "detail",
        ])
    return f, writer


def send_all_emails(sender_email, sender_password, subject, body_template, recipients, template_path):
    log_file, log_writer = init_log()
    sent_count = 0
    failed_count = 0

    context = ssl.create_default_context()
    try:
        with smtplib.SMTP_SSL(GMAIL_SMTP_HOST, GMAIL_SMTP_PORT, context=context, timeout=30) as server:
            server.login(sender_email, sender_password)

            for recipient in recipients:
                timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                email = recipient["email"]
                try:
                    msg = build_email(sender_email, recipient, subject, body_template, template_path)
                    server.send_message(msg)
                    print(f"[SENT] {recipient['trust_name']} -> {email}")
                    log_writer.writerow([
                        timestamp, sender_email, recipient["sno"], recipient["trust_name"],
                        email, subject, os.path.basename(template_path), "SENT", ""
                    ])
                    sent_count += 1
                except Exception as exc:
                    print(f"[FAILED] {recipient['trust_name']} -> {email}: {exc}")
                    log_writer.writerow([
                        timestamp, sender_email, recipient["sno"], recipient["trust_name"],
                        email, subject, os.path.basename(template_path), "FAILED", str(exc)
                    ])
                    failed_count += 1
                log_file.flush()
    finally:
        log_file.close()

    return sent_count, failed_count


def main():
    section("CLIENT EXCEL REQUEST EMAIL AUTOMATER")
    print("This program sends the SAME donation-details Excel template to every trust.")
    print()

    master_path, template_path = classify_excel_files()
    validate_template(template_path)
    recipients = load_recipients(master_path)

    sender_email, sender_password = get_authenticated_gmail_credentials()
    subject, body_template = prompt_email_configuration()

    if not delivery_confirmation(sender_email, recipients, template_path, subject):
        print("Cancelled. No emails were sent.")
        return

    sent_count, failed_count = send_all_emails(
        sender_email,
        sender_password,
        subject,
        body_template,
        recipients,
        template_path,
    )
    sender_password = None

    section("DELIVERY SUMMARY")
    print(f"Sent:   {sent_count}")
    print(f"Failed: {failed_count}")
    print(f"Log:    {os.path.abspath(LOG_FILE)}")
    print()
    print("The original client-input Excel was attached to every email.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelled by user.")
        sys.exit(1)
