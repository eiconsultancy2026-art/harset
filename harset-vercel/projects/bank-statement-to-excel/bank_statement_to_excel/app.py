import os
import re
from copy import copy
from datetime import datetime

import pdfplumber
import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_FILE = os.path.join(BASE_DIR, "AccountingVouchers(2).xlsx")

st.set_page_config(page_title="Bank Statement to Accounting Voucher", page_icon="🏦", layout="wide")


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value):
    value = clean_text(value)
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return None


def parse_amount(value):
    value = clean_text(value).replace(",", "")
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def extract_transactions(pdf_bytes):
    """Extract bank transaction rows from a tabular PDF statement."""
    transactions = []

    with pdfplumber.open(pdf_bytes) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            for table in page.extract_tables() or []:
                if not table:
                    continue

                for row in table:
                    if not row or len(row) < 7:
                        continue

                    sno_text = clean_text(row[0])
                    if not sno_text.isdigit():
                        continue

                    sno = int(sno_text)
                    # Axis statement's transaction total is not a transaction.
                    if sno == 172:
                        continue

                    tx_date = parse_date(row[1])
                    value_date = parse_date(row[2])
                    particulars = clean_text(row[3])
                    amount = parse_amount(row[4])
                    drcr = clean_text(row[5]).upper()
                    balance = parse_amount(row[6])
                    cheque = clean_text(row[7]) if len(row) > 7 else ""
                    branch = clean_text(row[8]) if len(row) > 8 else ""

                    if not tx_date or amount is None or drcr not in {"CR", "DR"}:
                        continue

                    transactions.append({
                        "sno": sno,
                        "transaction_date": tx_date,
                        "value_date": value_date,
                        "particulars": particulars,
                        "amount": amount,
                        "drcr": drcr,
                        "balance": balance,
                        "cheque": cheque,
                        "branch": branch,
                        "page": page_no,
                    })

    # Prevent duplicate rows when a PDF extraction engine returns the same table twice.
    seen = set()
    unique = []
    for tx in transactions:
        key = (tx["sno"], tx["transaction_date"], tx["amount"], tx["drcr"])
        if key not in seen:
            seen.add(key)
            unique.append(tx)

    return sorted(unique, key=lambda x: x["sno"])


def normalize_header(value):
    value = clean_text(value).lower()
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def find_template_columns(ws):
    """Find columns by the actual uploaded template headers."""
    cols = {}
    for col in range(1, ws.max_column + 1):
        header = normalize_header(ws.cell(1, col).value)
        if header == "voucher date":
            cols["date"] = col
        elif header == "voucher type name":
            cols["voucher_type"] = col
        elif header == "voucher number":
            cols["voucher_number"] = col
        elif header == "buyer supplier address":
            cols["address"] = col
        elif header == "buyer supplier pincode":
            cols["pincode"] = col
        elif header == "ledger name":
            cols["ledger_name"] = col
        elif header == "ledger amount":
            cols["amount"] = col
        elif header == "ledger amount dr cr":
            cols["drcr"] = col
        elif header == "item name":
            cols["item_name"] = col
        elif header == "billed quantity":
            cols["quantity"] = col
        elif header == "item rate":
            cols["item_rate"] = col
        elif header == "item rate per":
            cols["item_rate_per"] = col
        elif header == "item amount":
            cols["item_amount"] = col
        elif header == "change mode":
            cols["change_mode"] = col
    return cols


def copy_row_style(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def generate_excel(transactions, output_path):
    """Populate the original template without changing its header/read-me sheet."""
    wb = load_workbook(TEMPLATE_FILE)
    ws = wb["Accounting Voucher"] if "Accounting Voucher" in wb.sheetnames else wb.active

    # The supplied template currently has only the header row. If a future template
    # has a styled first data row, preserve it before clearing old data.
    style_row = 2 if ws.max_row >= 2 else None
    style_snapshot = None
    if style_row:
        style_snapshot = [copy(ws.cell(style_row, c)._style) for c in range(1, ws.max_column + 1)]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    cols = find_template_columns(ws)

    for excel_row, tx in enumerate(transactions, start=2):
        if style_snapshot:
            for col, style in enumerate(style_snapshot, start=1):
                ws.cell(excel_row, col)._style = copy(style)

        if "date" in cols:
            ws.cell(excel_row, cols["date"]).value = tx["transaction_date"]
            ws.cell(excel_row, cols["date"]).number_format = "dd-mm-yyyy"

        if "voucher_type" in cols:
            ws.cell(excel_row, cols["voucher_type"]).value = "Receipt" if tx["drcr"] == "CR" else "Payment"

        if "voucher_number" in cols:
            ws.cell(excel_row, cols["voucher_number"]).value = tx["sno"]

        # Bank particulars are the only reliable party/ledger text in this statement.
        if "ledger_name" in cols:
            ws.cell(excel_row, cols["ledger_name"]).value = tx["particulars"]

        if "amount" in cols:
            ws.cell(excel_row, cols["amount"]).value = tx["amount"]

        if "drcr" in cols:
            ws.cell(excel_row, cols["drcr"]).value = "Cr" if tx["drcr"] == "CR" else "Dr"

        # These fields are deliberately left blank because the bank statement
        # does not provide them in the required template format:
        # address, pincode, item name, quantity, item rate, item rate per,
        # item amount, change mode.

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(output_path)

    return cols


st.title("🏦 Bank Statement → Accounting Voucher Excel")
st.write("Upload a bank statement PDF. The app will populate the supplied AccountingVouchers template and leave unavailable fields blank.")

if not os.path.exists(TEMPLATE_FILE):
    st.error("AccountingVouchers(2).xlsx is missing from the application folder.")
    st.stop()

uploaded = st.file_uploader("Upload Bank Statement PDF", type=["pdf"])

if uploaded:
    st.success(f"Selected: {uploaded.name}")

    if st.button("Generate Excel", type="primary"):
        try:
            with st.spinner("Extracting transactions and creating Excel..."):
                transactions = extract_transactions(uploaded)

                if not transactions:
                    st.error("No transaction rows were detected. This version expects a text/table-based bank statement PDF.")
                    st.stop()

                base = os.path.splitext(uploaded.name)[0]
                output_path = os.path.join(BASE_DIR, f"{base}_AccountingVouchers.xlsx")
                columns = generate_excel(transactions, output_path)

            st.success(f"Done — {len(transactions)} transactions imported.")

            st.subheader("Imported fields")
            st.write({
                "Voucher Date": "Bank Transaction Date",
                "Voucher Type Name": "CR → Receipt, DR → Payment",
                "Voucher Number": "Bank S.No",
                "Ledger Name": "Bank Particulars",
                "Ledger Amount": "Bank Amount",
                "Ledger Amount Dr/Cr": "CR → Cr, DR → Dr",
                "Unavailable template fields": "Left blank",
            })

            with open(output_path, "rb") as f:
                st.download_button(
                    "⬇️ Download Excel",
                    data=f.read(),
                    file_name=os.path.basename(output_path),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

        except Exception as exc:
            st.error("The bank statement could not be processed.")
            st.exception(exc)
