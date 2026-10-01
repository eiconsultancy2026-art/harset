import io
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
import streamlit as st

st.set_page_config(page_title="Bank Statement → Tally", page_icon="🏦", layout="wide")

# ---------------- Parsing ----------------
DATE_PATTERNS = [
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y",
    "%d %b %Y", "%d %B %Y", "%Y-%m-%d"
]


def _normalise_name(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value).strip().lower()).strip()


def _normalise_columns(df):
    df = df.copy()
    df.columns = [_normalise_name(c) for c in df.columns]
    aliases = {
        "date": [
            "date", "transaction date", "txn date", "value date",
            "posting date", "tran date", "transaction dt", "txn dt"
        ],
        "narration": [
            "narration", "description", "details", "particulars",
            "transaction details", "remarks", "transaction description",
            "narrative", "reference", "transaction remarks"
        ],
        "debit": [
            "debit", "withdrawal", "withdrawals", "debit amount",
            "withdrawal amount", "dr", "withdrawal amt", "debit amt"
        ],
        "credit": [
            "credit", "deposit", "deposits", "credit amount",
            "deposit amount", "cr", "deposit amt", "credit amt"
        ],
        "amount": ["amount", "transaction amount", "txn amount", "amount inr"],
        "type": ["type", "transaction type", "dr cr", "credit debit", "cr dr"],
    }

    out = {}
    for target, names in aliases.items():
        for name in names:
            if name in df.columns:
                out[target] = df[name]
                break

    # Fuzzy fallback for bank-specific column names.
    if "date" not in out:
        for c in df.columns:
            if "date" in c or c.endswith(" dt"):
                out["date"] = df[c]
                break

    if "narration" not in out:
        for c in df.columns:
            if any(x in c for x in ["narr", "desc", "particular", "detail", "remark", "reference"]):
                out["narration"] = df[c]
                break

    result = pd.DataFrame(out)

    for c in ["debit", "credit", "amount"]:
        if c in result:
            s = result[c].astype(str).str.strip()
            # Handle accounting negatives such as (1,234.50), currency symbols,
            # commas, spaces and INR/₹ labels.
            s = s.str.replace(",", "", regex=False)
            s = s.str.replace("₹", "", regex=False)
            s = s.str.replace("INR", "", regex=False)
            s = s.str.replace(r"\((.*?)\)", r"-\1", regex=True)
            s = s.str.replace(r"[^0-9.\-]", "", regex=True)
            result[c] = pd.to_numeric(s, errors="coerce")

    if "date" in result:
        result["date"] = pd.to_datetime(result["date"], errors="coerce", dayfirst=True)

    return result


def _parse_dataframe(df):
    df = _normalise_columns(df)
    if df.empty or "date" not in df.columns:
        return pd.DataFrame()

    for c in ["debit", "credit", "amount"]:
        if c not in df.columns:
            df[c] = 0.0

    df["debit"] = df["debit"].fillna(0)
    df["credit"] = df["credit"].fillna(0)

    # Some bank files have one Amount column plus DR/CR or Type.
    if "amount" in df.columns:
        typ = df.get("type", pd.Series([""] * len(df), index=df.index)).astype(str).str.lower()
        missing = df["debit"].eq(0) & df["credit"].eq(0) & df["amount"].notna()
        debit_mask = missing & typ.str.contains(r"dr|debit|withdraw", regex=True, na=False)
        credit_mask = missing & ~typ.str.contains(r"dr|debit|withdraw", regex=True, na=False)
        df.loc[debit_mask, "debit"] = df.loc[debit_mask, "amount"].abs()
        df.loc[credit_mask, "credit"] = df.loc[credit_mask, "amount"].abs()

    df["narration"] = (
        df.get("narration", pd.Series([""] * len(df), index=df.index))
        .fillna("")
        .astype(str)
        .str.strip()
    )

    df = df[df["date"].notna() & ((df["debit"] != 0) | (df["credit"] != 0))].copy()
    if df.empty:
        return pd.DataFrame()

    df["debit"] = df["debit"].abs()
    df["credit"] = df["credit"].abs()
    return df[["date", "narration", "debit", "credit"]].reset_index(drop=True)


def _find_header_row(raw):
    """Find a likely bank transaction header anywhere in the first 50 rows."""
    if raw is None or raw.empty:
        return 0

    for i in range(min(50, len(raw))):
        values = [_normalise_name(x) for x in raw.iloc[i].tolist()]
        joined = " | ".join(values)
        has_date = any("date" in x for x in values)
        has_money = any(
            any(token in x for token in ["debit", "credit", "withdraw", "deposit", "amount", "dr", "cr"])
            for x in values
        )
        has_description = any(
            any(token in x for token in ["narr", "desc", "particular", "detail", "remark", "reference"])
            for x in values
        )
        if has_date and (has_money or has_description):
            return i
        # Some statements call the date field Txn Date and only use Dr/Cr.
        if "txn date" in joined and (" dr " in f" {joined} " or " cr " in f" {joined} "):
            return i
    return 0


def _read_excel_bytes(data, password=""):
    """Read normal, password-protected, and mislabeled Excel files."""
    last_errors = []
    working = data

    # Password-protected Office files are encrypted and are not valid ZIP/XLS
    # files until decrypted. msoffcrypto handles the common Office encryption.
    if password:
        try:
            import msoffcrypto
            source = io.BytesIO(data)
            office = msoffcrypto.OfficeFile(source)
            if office.is_encrypted():
                decrypted = io.BytesIO()
                office.load_key(password=password)
                office.decrypt(decrypted)
                working = decrypted.getvalue()
        except Exception as exc:
            last_errors.append(f"Password/decryption: {exc}")

    # Try modern .xlsx first.
    try:
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(working), read_only=True, data_only=True)
        frames = []
        for sheet in book.sheetnames:
            ws = book[sheet]
            raw = pd.DataFrame(ws.values)
            if raw.empty:
                continue
            header = _find_header_row(raw)
            for candidate in dict.fromkeys([header, 0, 1, 2, 3, 4]):
                if candidate >= len(raw):
                    continue
                frame = raw.iloc[candidate + 1:].copy()
                frame.columns = raw.iloc[candidate].tolist()
                parsed = _parse_dataframe(frame)
                if not parsed.empty:
                    frames.append(parsed)
                    break
        if frames:
            return pd.concat(frames, ignore_index=True), last_errors
    except Exception as exc:
        last_errors.append(f"openpyxl: {exc}")

    # Try legacy .xls / BIFF, including files with a wrong .xlsx extension.
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=working)
        frames = []
        for sheet_name in book.sheet_names():
            sheet = book.sheet_by_name(sheet_name)
            raw = pd.DataFrame([sheet.row_values(r) for r in range(sheet.nrows)])
            if raw.empty:
                continue
            header = _find_header_row(raw)
            for candidate in dict.fromkeys([header, 0, 1, 2, 3, 4]):
                if candidate >= len(raw):
                    continue
                frame = raw.iloc[candidate + 1:].copy()
                frame.columns = raw.iloc[candidate].tolist()
                parsed = _parse_dataframe(frame)
                if not parsed.empty:
                    frames.append(parsed)
                    break
        if frames:
            return pd.concat(frames, ignore_index=True), last_errors
    except Exception as exc:
        last_errors.append(f"xlrd: {exc}")

    return pd.DataFrame(), last_errors


def _read_csv_bytes(data):
    errors = []
    for encoding in ["utf-8-sig", "utf-8", "cp1252", "latin1"]:
        try:
            raw = pd.read_csv(io.BytesIO(data), header=None, encoding=encoding, sep=None, engine="python")
            header = _find_header_row(raw)
            for candidate in dict.fromkeys([header, 0, 1, 2, 3, 4]):
                if candidate >= len(raw):
                    continue
                frame = raw.iloc[candidate + 1:].copy()
                frame.columns = raw.iloc[candidate].tolist()
                parsed = _parse_dataframe(frame)
                if not parsed.empty:
                    return parsed, errors
        except Exception as exc:
            errors.append(f"CSV {encoding}: {exc}")
    return pd.DataFrame(), errors


def _read_pdf_bytes(data):
    errors = []
    try:
        import pdfplumber
        tables = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables() or []:
                    if not table or len(table) < 2:
                        continue
                    raw = pd.DataFrame(table)
                    header = _find_header_row(raw)
                    if header >= len(raw) - 1:
                        continue
                    frame = raw.iloc[header + 1:].copy()
                    frame.columns = raw.iloc[header].tolist()
                    parsed = _parse_dataframe(frame)
                    if not parsed.empty:
                        tables.append(parsed)
        if tables:
            return pd.concat(tables, ignore_index=True), errors
    except Exception as exc:
        errors.append(f"PDF table extraction: {exc}")
    return pd.DataFrame(), errors


def read_uploaded(upload, password=""):
    suffix = Path(upload.name).suffix.lower()
    data = upload.getvalue()

    if suffix in [".xlsx", ".xls"]:
        return _read_excel_bytes(data, password=password)
    if suffix == ".csv":
        return _read_csv_bytes(data)
    if suffix == ".pdf":
        return _read_pdf_bytes(data)
    return pd.DataFrame(), [f"Unsupported file type: {suffix}"]


# ---------------- Tally exports ----------------
def make_tally_import_dataframe(df):
    """Build the exact seven-column layout used by the supplied Tally bank import template."""
    out = pd.DataFrame()
    out["Date"] = pd.to_datetime(df["date"], errors="coerce").dt.strftime("%d %b %Y")
    out["Narration"] = df["narration"].fillna("").astype(str).str.strip()
    # The source statement parser does not invent cheque/reference numbers.
    # Keep this column blank unless the source explicitly supplies one in a future parser extension.
    out["Chq./Ref.No."] = ""
    out["Value Dt"] = pd.to_datetime(df["date"], errors="coerce").dt.strftime("%d %b %Y")
    debit = pd.to_numeric(df["debit"], errors="coerce").fillna(0).abs().round(2)
    credit = pd.to_numeric(df["credit"], errors="coerce").fillna(0).abs().round(2)
    # The supplied Tally template uses '-' for the side with no transaction amount.
    out["Withdrawal Amt."] = debit.map(lambda x: "-" if x == 0 else float(x))
    out["Deposit Amt."] = credit.map(lambda x: "-" if x == 0 else float(x))
    # Closing balance is intentionally blank when the source statement does not provide it.
    # It must not be fabricated from an unknown opening balance.
    out["Closing Balance"] = ""
    return out


def _xlsx_bytes_to_xls(xlsx_bytes):
    """Convert the generated XLSX to legacy XLS when LibreOffice/soffice is installed."""
    candidates = [shutil.which("soffice"), shutil.which("libreoffice")]
    if not any(candidates):
        return None, "LibreOffice/soffice was not found; XLSX was generated instead."

    exe = next(c for c in candidates if c)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "tally_import.xlsx"
        outdir = Path(td) / "out"
        outdir.mkdir()
        src.write_bytes(xlsx_bytes)
        proc = subprocess.run(
            [exe, "--headless", "--convert-to", "xls", "--outdir", str(outdir), str(src)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=60,
        )
        target = outdir / "tally_import.xls"
        if proc.returncode == 0 and target.exists():
            return target.read_bytes(), None
        return None, (proc.stderr or proc.stdout or "LibreOffice could not create XLS.").strip()


def make_tally_import_workbook(df):
    """Return XLSX bytes matching the supplied Tally import template layout."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    data = make_tally_import_dataframe(df)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet 1"

    # These three rows mirror the supplied Tally import template.
    ws.cell(1, 1).value = "*" * 188
    headers = ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]
    separators = ["********", "**********************************", "************", "********", "******************", "******************", "******************"]
    for col, value in enumerate(headers, 1):
        ws.cell(2, col).value = value
    for col, value in enumerate(separators, 1):
        ws.cell(3, col).value = value

    for row_idx, row in enumerate(data.itertuples(index=False), 4):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row_idx, col_idx).value = value

    widths = [10.13, 104.7, 18.7, 10.99, 18.0, 18.7, 18.0]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")

    for cell in ws[2]:
        cell.font = Font(name="Arial", size=10, bold=True)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A2:G{max(3, ws.max_row)}"

    # Keep dates as text in the exact display style shown by the supplied template.
    # Amounts are numeric when present and '-' when the side is empty.
    for r in range(4, ws.max_row + 1):
        for c in [5, 6, 7]:
            ws.cell(r, c).number_format = '#,##0.00'

    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_tally_xml(df, bank_ledger, contra_ledger):
    root = ET.Element("ENVELOPE")
    header = ET.SubElement(root, "HEADER")
    ET.SubElement(header, "VERSION").text = "1"
    ET.SubElement(header, "TALLYREQUEST").text = "Import"
    ET.SubElement(header, "TYPE").text = "Data"
    body = ET.SubElement(root, "BODY")
    imp = ET.SubElement(body, "IMPORTDATA")
    desc = ET.SubElement(imp, "REQUESTDESC")
    ET.SubElement(desc, "REPORTNAME").text = "Vouchers"
    req = ET.SubElement(imp, "REQUESTDATA")

    for _, r in df.iterrows():
        vtype = "Payment" if r["debit"] > 0 else "Receipt"
        msg = ET.SubElement(req, "TALLYMESSAGE")
        voucher = ET.SubElement(msg, "VOUCHER", {"VCHTYPE": vtype, "ACTION": "Create"})
        ET.SubElement(voucher, "DATE").text = r["date"].strftime("%Y%m%d")
        ET.SubElement(voucher, "VOUCHERTYPENAME").text = vtype
        ET.SubElement(voucher, "NARRATION").text = str(r["narration"])

        bank_amt = -float(r["debit"]) if r["debit"] > 0 else float(r["credit"])
        party_amt = -bank_amt

        e1 = ET.SubElement(voucher, "ALLLEDGERENTRIES.LIST")
        ET.SubElement(e1, "LEDGERNAME").text = bank_ledger
        ET.SubElement(e1, "ISDEEMEDPOSITIVE").text = "Yes" if bank_amt < 0 else "No"
        ET.SubElement(e1, "AMOUNT").text = f"{bank_amt:.2f}"

        e2 = ET.SubElement(voucher, "ALLLEDGERENTRIES.LIST")
        ET.SubElement(e2, "LEDGERNAME").text = contra_ledger
        ET.SubElement(e2, "ISDEEMEDPOSITIVE").text = "Yes" if party_amt < 0 else "No"
        ET.SubElement(e2, "AMOUNT").text = f"{party_amt:.2f}"
        ET.SubElement(voucher, "EFFECTIVEDATE").text = r["date"].strftime("%Y%m%d")

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


# ---------------- UI ----------------
st.title("🏦 Bank Statement → Tally Ready Converter")
st.caption("Upload a bank statement, review the parsed transactions, then export Excel or Tally XML.")

with st.sidebar:
    st.header("Tally Settings")
    bank_ledger = st.text_input("Bank Ledger Name", "Bank Account")
    contra_ledger = st.text_input("Default Contra Ledger", "Suspense")
    excel_password = st.text_input(
        "Excel Password",
        value="",
        type="password",
        placeholder="Enter password if the Excel file is protected",
        help="Used only to decrypt a password-protected Excel statement before reading it."
    )
    st.divider()
    st.info("For accounting safety, transactions are shown for review before export. The app does not invent unknown ledgers.")

upload = st.file_uploader("Upload Bank Statement", type=["xlsx", "xls", "csv", "pdf"])

if upload:
    with st.spinner("Reading statement..."):
        df, parse_errors = read_uploaded(upload, password=excel_password)

    if df.empty:
        st.error("Could not read transactions from this statement.")
        if parse_errors:
            # Keep technical details compact; the important part is the actionable message.
            with st.expander("Technical details"):
                for err in parse_errors[-5:]:
                    st.code(str(err))
        st.info(
            "For password-protected Excel, enter the Excel Password in the left sidebar. "
            "The app supports normal .xlsx/.xls files, password-protected Office Excel files, "
            "and bank files whose .xlsx extension is actually a legacy .xls workbook."
        )
        st.stop()

    st.success(f"Detected {len(df)} transactions.")
    st.subheader("Review & Edit")
    edited = st.data_editor(
        df,
        use_container_width=True,
        num_rows="dynamic",
        column_config={
            "date": st.column_config.DateColumn("Date"),
            "debit": st.column_config.NumberColumn("Debit", format="₹ %.2f"),
            "credit": st.column_config.NumberColumn("Credit", format="₹ %.2f"),
        },
        hide_index=True,
    )

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Transactions", len(edited))
    c2.metric("Total Debit", f"₹{edited.debit.sum():,.2f}")
    c3.metric("Total Credit", f"₹{edited.credit.sum():,.2f}")
    c4.metric("Net", f"₹{(edited.credit.sum() - edited.debit.sum()):,.2f}")

    tally = make_tally_import_dataframe(edited)
    st.subheader("Tally Bank Import Preview")
    st.dataframe(tally, use_container_width=True, hide_index=True)

    st.divider()
    st.subheader("Export")
    col1, col2, col3 = st.columns(3)
    import_xlsx = make_tally_import_workbook(edited)
    with col1:
        st.download_button(
            "⬇️ Download Tally Import XLSX",
            import_xlsx,
            "Bank_Import_Tally.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with col2:
        import_xls, xls_error = _xlsx_bytes_to_xls(import_xlsx)
        if import_xls:
            st.download_button(
                "⬇️ Download Tally Import XLS",
                import_xls,
                "Bank_Import_Tally.xls",
                "application/vnd.ms-excel",
                use_container_width=True,
            )
        else:
            st.button("XLS requires LibreOffice", disabled=True, use_container_width=True)
            if xls_error:
                st.caption(xls_error)
    with col3:
        xml = make_tally_xml(edited, bank_ledger, contra_ledger)
        st.download_button(
            "⬇️ Download Tally XML",
            xml,
            "tally_vouchers.xml",
            "application/xml",
            use_container_width=True,
        )

    st.warning(
        "The Excel export now follows the supplied Tally Bank Import template exactly: "
        "Date, Narration, Chq./Ref.No., Value Dt, Withdrawal Amt., Deposit Amt., Closing Balance. "
        "Closing Balance is left blank unless it is available from the source statement; the app does not invent it."
    )
else:
    st.info("Start by uploading an Excel, CSV, or PDF bank statement.")
    st.markdown("""
### Supported workflow
**Bank statement → automatic column detection → transaction review → Tally Excel/XML**

For best accuracy, upload the original **Excel/CSV downloaded from your bank** rather than a PDF converted to Excel.
""")
